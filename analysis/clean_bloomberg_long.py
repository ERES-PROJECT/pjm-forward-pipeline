"""Convert Bloomberg FPWPM/FPWOM monthly contract wide files into a matched long table.
Run from cost-pipeline root: python analysis/clean_bloomberg_long.py
Requires: pip install openpyxl
"""
from pathlib import Path
from datetime import datetime, date
from collections import Counter
import csv, re, math
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
OUT = ROOT / 'analysis' / 'output'
OUT.mkdir(parents=True, exist_ok=True)
MONTH = {c:i for i,c in enumerate('FGHJKMNQUVXZ', 1)}
PATTERN = re.compile(r'\b(FPWPM|FPWOM)\s+([FGHJKMNQUVXZ])(\d{2})\b')

def read_book(path, prefix):
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    columns = {}
    for i, h in enumerate(header[1:], 1):
        if h is None or str(h).strip() == '':
            continue
        m = PATTERN.search(str(h))
        if not m or m.group(1) != prefix:
            raise ValueError(f'Unexpected column in {path.name}: {h!r}')
        code = m.group(2) + m.group(3)
        if code in columns.values():
            raise ValueError(f'Duplicate contract column: {code}')
        columns[i] = code
    result = {}
    invalid = []
    for rownum, row in enumerate(rows, 2):
        dt = row[0]
        if dt is None:
            if any(v is not None for v in row[1:]):
                invalid.append((rownum, 'missing date'))
            continue
        if isinstance(dt, datetime): dt = dt.date()
        elif not isinstance(dt, date):
            invalid.append((rownum, f'invalid date {dt!r}'))
            continue
        for i, code in columns.items():
            raw = row[i] if i < len(row) else None
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            try:
                value = float(raw)
                if not math.isfinite(value): raise ValueError()
            except (TypeError, ValueError):
                invalid.append((rownum, f'non-numeric {code}: {raw!r}'))
                continue
            key = (dt, code)
            if key in result:
                raise ValueError(f'Duplicate observation-date/contract in {path.name}: {key}')
            result[key] = value
    if invalid:
        raise ValueError(f'{path.name}: {len(invalid)} invalid cells; examples: {invalid[:8]}')
    return result, set(columns.values())

def main():
    on, oncodes = read_book(DATA / 'OnPeak.xlsx', 'FPWPM')
    off, offcodes = read_book(DATA / 'OffPeak.xlsx', 'FPWOM')
    if oncodes != offcodes:
        print('WARNING: Contract sets differ:', sorted(oncodes ^ offcodes))
    keys = sorted(set(on) | set(off), key=lambda k:(k[0], 2000+int(k[1][1:]), MONTH[k[1][0]]))
    output = OUT / 'bloomberg_on_off_peak_clean_long.csv'
    counts = Counter()
    with output.open('w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['Observation_Date','Delivery_Month','Contract_Code','OnPeak_Price','OffPeak_Price','Price_Unit','Price_Status'])
        for dt, code in keys:
            delivery = date(2000+int(code[1:]), MONTH[code[0]], 1)
            a, b = on.get((dt,code)), off.get((dt,code))
            status = 'BOTH' if a is not None and b is not None else ('ON_ONLY' if a is not None else 'OFF_ONLY')
            counts[status] += 1
            writer.writerow([dt.isoformat(),delivery.isoformat(),code,'' if a is None else a,'' if b is None else b,'USD/MWh',status])
    print('Output:', output)
    print('Unique matched keys:', len(keys), 'On values:',len(on),'Off values:',len(off))
    print('Statuses:',dict(counts))
    print('Date range:', keys[0][0] if keys else None, 'to', keys[-1][0] if keys else None)
    print('No interpolation, forward fill, or zero fill performed.')

if __name__ == '__main__': main()
